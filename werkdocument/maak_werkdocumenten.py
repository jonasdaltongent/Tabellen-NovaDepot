#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het werkdocument voor les 05 "Alles op een rij" (Tekstverwerking 4: tabellen,
NovaDepot, ORLO).

  TV4_Tabellen.docx   het werkdocument dat de leerling INLEVERT
                      deel A: een tabel met de leveringen van dinsdag 13 oktober (rij 1 ingevuld)
                      deel B: plaats voor een zelfgemaakte deelnemerslijst
                      vragen, zelfcontrole en een optionele uitbreiding

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx

LET OP bij aanpassen:
 - Deel A: de tabel heeft 4 even brede kolommen. Dat is OPZETTELIJK: lange namen zoals
   "Koffiebranderij De Bonenbaas" passen dan niet op één regel, zodat de leerling in stap 4
   een kolom breder sleept. De koprij is NIET opgemaakt: dat doet de leerling (stap 4).
 - Rij 1 van deel A is al ingevuld: dat is de demo van de leraar.
 - Het briefje en de wijzigingen moeten kloppen met de verbetersleutel in README.md §4:
   na stap 3 staan er zes leveringen, in volgorde van het uur, zonder Frisdranken Bubbel,
   met Tuincentrum Groenvinger om 11.15 uur, en een lege kolom Afgetekend.
 - Het document zelf is wel correct opgemaakt (A4, staand, marges 2 cm, koptekst en
   paginanummer): dat is wat de leerling in les 04 leerde.
 - NovaDepot, alle leveranciers en alle namen van jobstudenten zijn fictief.
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HERE, "TV4_Tabellen.docx")

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x4D, 0x55, 0x73)
GRIJS = RGBColor(0x80, 0x80, 0x80)

# Het briefje van het onthaal (deel A). Volgorde = volgorde van het uur.
LEVERINGEN = [
    ("7.30", "Bakkerij Korstjes", "2", "3"),              # rij 1 = de demo, al ingevuld
    ("8.15", "Papierhandel Vellekens", "4", "4"),
    ("9.00", "Frisdranken Bubbel", "6", "3"),             # valt weg in stap 3
    ("10.30", "Koffiebranderij De Bonenbaas", "1", "4"),
    ("13.00", "Speelgoed Tolletje", "3", "3"),
    ("14.45", "Schoonmaak Glanzmann", "2", "4"),
]
EXTRA_LEVERING = ("11.15", "Tuincentrum Groenvinger", "2", "3")   # komt erbij in stap 3

# Het mailtje van de personeelsdienst (deel B)
JOBSTUDENTEN = [
    ("Lina Haddad", "magazijn"),
    ("Emre Yilmaz", "magazijn"),
    ("Noor Verbeke", "onthaal"),
    ("Kobe De Smet", "verzending"),
    ("Ayla Janssens", "onthaal"),
]


# ---------- hulpfuncties ----------
def lettertype(run, naam="Arial"):
    run.font.name = naam
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), naam)


def basis_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "Arial")
    taal = OxmlElement("w:lang")
    taal.set(qn("w:val"), "nl-BE")
    rpr.append(taal)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.1

    # Correct opgemaakt, zoals in les 04: A4 staand, marges 2 cm, koptekst en paginanummer.
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    s.header_distance = s.footer_distance = Cm(1)

    kop = s.header.paragraphs[0]
    kop.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = kop.add_run("NovaDepot · Toegepaste Informatica")
    r.font.size = Pt(9)
    r.font.color.rgb = GRIJS

    voet = s.footer.paragraphs[0]
    voet.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = voet.add_run("Pagina ")
    r.font.size = Pt(9)
    r.font.color.rgb = GRIJS
    veld = OxmlElement("w:fldSimple")
    veld.set(qn("w:instr"), "PAGE")
    vr = OxmlElement("w:r")
    vrpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    vrpr.append(sz)
    vr.append(vrpr)
    vt = OxmlElement("w:t")
    vt.text = "1"
    vr.append(vt)
    veld.append(vr)
    voet._p.append(veld)

    doc.core_properties.title = "TV4 Tabellen - werkdocument NovaDepot"
    doc.core_properties.author = "Toegepaste Informatica"
    return doc


def tekst(doc, s, vet=False, klein=False, cursief=False, na=6, grootte=11, kleur=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(9.5 if klein else grootte)
    if klein:
        r.font.color.rgb = GREY
    if kleur is not None:
        r.font.color.rgb = kleur
    return p


def kop(doc, s, grootte=13, voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(voor)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(s)
    lettertype(r)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def notitie(doc, regels, titel):
    """Een briefje: alinea's met een lichte achtergrond (geen tabel, zodat het niet verwart)."""
    def arceer(p):
        ppr = p._p.get_or_add_pPr()
        rand = OxmlElement("w:pBdr")
        links = OxmlElement("w:left")
        links.set(qn("w:val"), "single")
        links.set(qn("w:sz"), "18")
        links.set(qn("w:space"), "6")
        links.set(qn("w:color"), "E0B040")
        rand.append(links)
        volg = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr")
        _voeg_in_voor(ppr, rand, volg)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "FFF7E3")
        _voeg_in_voor(ppr, shd, volg[1:])
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "170")
        ind.set(qn("w:right"), "170")
        _voeg_in_voor(ppr, ind, ("w:jc", "w:rPr"))

    p = tekst(doc, titel, vet=True, na=0)
    arceer(p)
    for regel in regels:
        p = tekst(doc, regel, cursief=True, na=0)
        arceer(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------- tabellen met vaste kolommen en randen (schemavolgorde gerespecteerd) ----------
def _voeg_in_voor(ouder, el, opvolgers):
    """Voeg el in vóór het eerste bestaande element uit opvolgers (OOXML-volgorde), anders achteraan."""
    for tag in opvolgers:
        ref = ouder.find(qn(tag))
        if ref is not None:
            ref.addprevious(el)
            return el
    ouder.append(el)
    return el


def vaste_tabel(t, breedtes_cm, randkleur="8C8C8C"):
    """Vaste kolombreedtes en dunne randen, zodat Google Documenten de tabel toont zoals bedoeld."""
    tbl = t._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _voeg_in_voor(tblpr, OxmlElement("w:tblW"),
                             ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                              "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    tblw.set(qn("w:w"), str(int(round(sum(breedtes_cm) * 567))))
    tblw.set(qn("w:type"), "dxa")
    randen = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + kant)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), randkleur)
        randen.append(el)
    _voeg_in_voor(tblpr, randen, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    indeling = OxmlElement("w:tblLayout")
    indeling.set(qn("w:type"), "fixed")
    _voeg_in_voor(tblpr, indeling, ("w:tblCellMar", "w:tblLook"))
    for kolom, b in zip(tbl.tblGrid.findall(qn("w:gridCol")), breedtes_cm):
        kolom.set(qn("w:w"), str(int(round(b * 567))))
    for rij in t.rows:
        for cel, b in zip(rij.cells, breedtes_cm):
            cel.width = Cm(b)
    return t


def schaduw(cel, kleur="E4E8F6"):
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcpr.append(shd)


def cel_tekst(cel, s, vet=False, grootte=11, cursief=False):
    cel.text = ""
    p = cel.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(grootte)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("_" * 78)
        lettertype(r)
        r.font.color.rgb = GRIJS


# ---------- het werkdocument ----------
def werkdocument():
    doc = basis_document()

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WERKDOCUMENT — ALLES OP EEN RIJ")
    lettertype(r)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 05 — Tabellen  ·  Tekstverwerking 4  ·  Toegepaste Informatica  ·  NovaDepot",
          klein=True, na=8)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    vaste_tabel(t, [7.0, 4.5, 5.5])
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        cel_tekst(t.rows[0].cells[i], label + " ", vet=True)

    kop(doc, "Zo werk je", 12, voor=10)
    for s in [
        "1.  Links op je scherm staat de lespagina. Daar lees je wat je moet doen.",
        "2.  In dit werkdocument maak je twee tabellen: deel A en deel B.",
        "3.  Onderaan beantwoord je vier vragen.",
        "4.  Alleen DIT document lever je in via Google Classroom.",
    ]:
        tekst(doc, s, na=1)

    # ---- deel A ----
    kop(doc, "Deel A — Leveringen van dinsdag 13 oktober")
    tekst(doc, "Tom van het magazijn wil de leveringen in één overzicht, om aan de muur van de "
               "ontvangstzone te hangen. Vul de tabel in met het briefje van het onthaal. "
               "Rij 1 deed je leraar voor.", klein=True)
    notitie(doc, ["Dinsdag 13 oktober verwachten we zes leveringen."] +
            ["Om {} uur brengt {} {} {} naar poort {}.".format(
                uur, lev, pal, "pallet" if pal == "1" else "pallets", poort)
             for uur, lev, pal, poort in LEVERINGEN],
            "Het briefje van het onthaal")

    tekst(doc, "Leveringen van dinsdag 13 oktober", vet=True, grootte=13, na=4)
    t = doc.add_table(rows=len(LEVERINGEN) + 1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    vaste_tabel(t, [4.25, 4.25, 4.25, 4.25])        # even breed: lange namen passen niet (opzettelijk)
    for i, kopje in enumerate(["Uur", "Leverancier", "Pallets", "Poort"]):
        cel_tekst(t.rows[0].cells[i], kopje)          # koprij bewust niet opgemaakt
    for j, waarde in enumerate(LEVERINGEN[0]):
        cel_tekst(t.rows[1].cells[j], waarde)         # rij 1 = de demo
    for i in range(2, len(LEVERINGEN) + 1):
        for j in range(4):
            cel_tekst(t.rows[i].cells[j], "")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    uur, lev, pal, poort = EXTRA_LEVERING
    notitie(doc, [
        "1. Frisdranken Bubbel komt niet: hun vrachtwagen is stuk.",
        "2. Er komt een levering bij: om {} uur brengt {} {} pallets naar poort {}.".format(uur, lev, pal, poort),
        "3. Tom wil zien wie elke levering controleerde. Hij wil rechts een kolom Afgetekend. "
        "Die blijft leeg: Tom tekent op papier.",
    ], "Om 8.00 uur belt het onthaal: er zijn drie wijzigingen")

    # ---- deel B ----
    kop(doc, "Deel B — Deelnemerslijst voor de infosessie")
    tekst(doc, "Op de eerste werkdag van de jobstudenten is er om 8.00 uur een infosessie in de "
               "kantine. Lotte van de personeelsdienst mailde de namen. Maak zelf een tabel: je "
               "gebruikt alleen de tabelkaart op de lespagina.", klein=True)
    notitie(doc, ["Hallo! Deze vijf jobstudenten komen naar de infosessie:"] +
            ["{} – {}".format(naam, afd) for naam, afd in JOBSTUDENTEN] +
            ["Groetjes, Lotte (personeelsdienst)"],
            "Het mailtje van Lotte")
    tekst(doc, "▶  Maak hier je deelnemerslijst: eerst de titel, daaronder de tabel.", vet=True,
          kleur=NAVY, na=2)
    doc.add_paragraph()
    doc.add_paragraph()

    # ---- vragen ----
    kop(doc, "Vragen")
    tekst(doc, "Vraag 1. Hoeveel pallets gaan er naar poort 4? Vond je dat sneller in het briefje "
               "of in je tabel? Waarom?", vet=True)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 2. Waarom maak je de koprij vet en geef je ze een kleur?", vet=True)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 3. Waarom maak je een lijst met kolommen als tabel, en niet met spaties of de "
               "Tab-toets?", vet=True)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 4. Welke stap vond je vandaag het moeilijkst? Waarom?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- zelfcontrole ----
    kop(doc, "Zelfcontrole — aankruisen vóór je inlevert")
    for s in [
        "In deel A staan zes leveringen, van 7.30 uur tot 14.45 uur, op volgorde van het uur.",
        "Frisdranken Bubbel staat er niet meer in. Tuincentrum Groenvinger staat om 11.15 uur.",
        "Rechts staat een kolom Afgetekend. Er zijn geen lege rijen.",
        "De koprij is vet en lichtblauw.",
        "De getallen staan in het midden van hun cel. Elke naam past op één regel.",
        "In deel B staat een titel en daaronder een tabel met drie kolommen en zes rijen.",
        "Mijn deelnemerslijst volgt de tabelkaart. De kolom Handtekening is leeg.",
        "De vier vragen zijn ingevuld.",
    ]:
        tekst(doc, "☐  " + s, na=1)

    # ---- extra ----
    kop(doc, "Extra — niet verplicht")
    tekst(doc, "Alleen als je al ingeleverd hebt. Klik in Classroom op Inleveren ongedaan maken. "
               "Lever daarna opnieuw in.", klein=True)
    tekst(doc, "a) Zet je deelnemerslijst op alfabetische volgorde. Kijk op de lespagina bij Extra.", vet=True, na=2)
    tekst(doc, "b) Geef de jobstudenten meer plaats om te tekenen: maak hun rijen hoger.", vet=True, na=2)

    doc.save(UIT)


if __name__ == "__main__":
    werkdocument()
    print("Klaar: " + os.path.basename(UIT))
    print("  deel A: tabel {} x 4, rij 1 ingevuld (de demo), koprij niet opgemaakt".format(len(LEVERINGEN) + 1))
    print("  deel B: plaats voor de deelnemerslijst ({} jobstudenten)".format(len(JOBSTUDENTEN)))
